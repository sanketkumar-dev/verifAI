# VerifAI - website ke routes. AI, database aur PDF ka kaam alag files mein hai.
import io, os, re, uuid, sqlite3
from datetime import datetime
from functools import wraps
from flask import Flask, request, jsonify, send_file, send_from_directory, session, redirect, url_for
from docx import Document
from PIL import Image
from pypdf import PdfReader
from werkzeug.security import generate_password_hash, check_password_hash   # password ko code (hash) mein badalta hai
from dotenv import load_dotenv

import database, ai_service, pdf_service

load_dotenv()
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY") or os.urandom(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
BASE = os.path.dirname(os.path.abspath(__file__))
DOCUMENT_TYPES = {".pdf", ".txt", ".docx", ".jpg", ".jpeg", ".png"}
DOCUMENT_MAX_SIZE = 10 * 1024 * 1024
database.create_tables()


def current_user_id():
    """Abhi kaun login hai (nahi to None)."""
    return session.get("user_id")


def need_login():
    return jsonify({"error": "Pehle login karo."}), 401


def login_required(function):
    @wraps(function)
    def check_login(*args, **kwargs):
        if not current_user_id():
            return need_login()
        return function(*args, **kwargs)
    return check_login


# ---------- 1. PAGES ----------
@app.route("/")
def home():
    # QR/link (/?id=...) bina login ke bhi khulna chahiye, baaki page ke liye login zaroori
    if not current_user_id() and not request.args.get("id"):
        return redirect(url_for("login_page", next=request.full_path))
    return send_from_directory(BASE, "index.html")

@app.route("/index.html")
def dashboard_page():
    return home()

@app.route("/login")
def login_page():
    if current_user_id():
        return redirect(url_for("home"))
    return send_from_directory(BASE, "login.html")

@app.route("/signup")
def signup_page():
    if current_user_id():
        return redirect(url_for("home"))
    return send_from_directory(BASE, "signup.html")

@app.route("/style.css")
def css():
    return send_from_directory(BASE, "style.css")

@app.route("/auth.css")
def auth_css():
    return send_from_directory(BASE, "auth.css")

@app.route("/icon.svg")
def icon():
    return send_from_directory(BASE, "icon.svg")

@app.route("/script.js")
def js():
    return send_from_directory(BASE, "script.js")


# ---------- 2. LOGIN / SIGNUP ----------
@app.route("/api/signup", methods=["POST"])
def signup():
    body = request.get_json(silent=True) or {}
    name = str(body.get("name", "")).strip()
    email = str(body.get("email", "")).strip().lower()
    password = str(body.get("password", ""))
    confirm_password = str(body.get("confirm_password", ""))
    if not name or len(name) > 100 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return jsonify({"error": "Sahi naam aur email address daalo."}), 400
    if len(password) < 6:
        return jsonify({"error": "Password kam se kam 6 akshar ka hona chahiye."}), 400
    if password != confirm_password:
        return jsonify({"error": "Dono password ek jaise hone chahiye."}), 400
    if database.find_user(email):
        return jsonify({"error": "Ye email pehle se bana hua hai. Login karo."}), 400
    try:
        database.add_user(name, email, generate_password_hash(password))
    except sqlite3.IntegrityError:
        return jsonify({"error": "Ye email pehle se bana hua hai. Login karo."}), 400
    return jsonify({"ok": True})

@app.route("/api/login", methods=["POST"])
def login():
    body = request.get_json(silent=True) or {}
    email = str(body.get("email", "")).strip().lower()
    user = database.find_user(email)
    if not user or not check_password_hash(user["password_hash"], str(body.get("password", ""))):
        return jsonify({"error": "Email ya password galat hai."}), 400
    session.clear()
    session["user_id"] = user["id"]
    return jsonify({"email": email})

@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})

@app.route("/api/me")
def me():
    user = database.find_user_by_id(current_user_id()) if current_user_id() else None
    return jsonify({"email": user["email"] if user else None})


# ---------- 3. JAANCH (check) ----------
@app.route("/api/verify", methods=["POST"])
@login_required
def verify():
    lang = request.form.get("lang", "en")
    mode = request.form.get("mode", "text")
    text = request.form.get("text", "").strip()
    report_content = text
    upload = request.files.get("file")
    content_type = request.form.get("content_type", "auto")
    try:
        if not os.getenv("GEMINI_API_KEY"):
            raise ValueError(".env file mein GEMINI_API_KEY nahi mili.")
        if content_type not in ("auto", "news", "link", "email"):
            raise ValueError("Choose a valid content type.")
        if mode not in ("text", "file", "document"):
            raise ValueError("Choose a valid verification type.")

        file_bytes = None
        analysis_kind = "Text"
        report_kind = "Text"
        mime = None
        if mode == "file":
            if not upload or not upload.filename:
                raise ValueError("Choose a photo or video first.")
            file_bytes = upload.read(14_000_001)
            if len(file_bytes) > 14_000_000:
                raise ValueError("Keep the photo or video smaller than 14 MB.")
            analysis_kind = "Video" if upload.mimetype.startswith("video/") else "Image"
            report_kind = analysis_kind
            mime = upload.mimetype
        elif mode == "document":
            if not upload or not upload.filename:
                raise ValueError("Choose a document first.")
            extension = os.path.splitext(upload.filename)[1].lower()
            if extension not in DOCUMENT_TYPES:
                raise ValueError("Unsupported file. Choose PDF, TXT, DOCX, JPG, JPEG or PNG.")
            file_bytes = upload.read(DOCUMENT_MAX_SIZE + 1)
            if len(file_bytes) > DOCUMENT_MAX_SIZE:
                raise ValueError("Documents must be 10 MB or smaller.")
            if not file_bytes:
                raise ValueError("The selected file is empty or unreadable.")
            report_kind = "Image" if extension in (".jpg", ".jpeg", ".png") else f"{extension[1:].upper()} Document"
            mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}.get(extension)
            if mime:
                expected_format = "JPEG" if extension in (".jpg", ".jpeg") else "PNG"
                try:
                    with Image.open(io.BytesIO(file_bytes)) as image:
                        if image.format != expected_format:
                            raise ValueError("Image extension does not match its file type.")
                        image.verify()
                except ValueError:
                    raise
                except Exception as error:
                    raise ValueError("This image is damaged or unreadable.") from error
            if extension == ".txt":
                try:
                    text = file_bytes.decode("utf-8-sig")
                except UnicodeDecodeError:
                    try:
                        text = file_bytes.decode("cp1252")
                    except UnicodeDecodeError as error:
                        raise ValueError("This TXT file is not readable text.") from error
            elif extension == ".pdf":
                try:
                    pdf_reader = PdfReader(io.BytesIO(file_bytes))
                    if pdf_reader.is_encrypted and not pdf_reader.decrypt(""):
                        raise ValueError("This PDF is password-protected and cannot be read.")
                    text = "\n".join(page.extract_text() or "" for page in pdf_reader.pages).strip()
                except ValueError:
                    raise
                except Exception as error:
                    raise ValueError("This PDF is damaged or could not be read.") from error
                if not text:
                    raise ValueError("No readable text was found in the PDF.")
            elif extension == ".docx":
                try:
                    word_doc = Document(io.BytesIO(file_bytes))
                    text = "\n".join(
                        [paragraph.text for paragraph in word_doc.paragraphs]
                        + [cell.text for table in word_doc.tables for row in table.rows for cell in row.cells]
                    ).strip()
                except Exception as error:
                    raise ValueError("This DOCX file is damaged or could not be read.") from error
                if not text:
                    raise ValueError("No readable text was found in the DOCX file.")

            if extension in (".jpg", ".jpeg", ".png"):
                analysis_kind = "Email" if content_type == "email" else "Image"
            else:
                analysis_kind = "Email" if content_type == "email" else "Link" if content_type == "link" else "Text"
        else:
            if not text:
                raise ValueError("Enter some content to check.")
            if content_type == "email":
                analysis_kind = "Email"
            elif content_type == "link" or (content_type == "auto" and re.match(r"https?://", text)):
                if not re.match(r"https?://", text):
                    raise ValueError("Enter a valid website address starting with http:// or https://.")
                analysis_kind = "Link"
            else:
                analysis_kind = "Text"

        if mode == "text":
            report_kind = analysis_kind
        elif mode == "document":
            report_content = text

        if mode == "document" and analysis_kind == "Link":
            if not re.fullmatch(r"https?://\S+", text):
                raise ValueError("For Website Link, the document must contain a single web address.")
            report_content = text
            text = ai_service.get_page_text(text)
        elif mode == "text" and analysis_kind == "Link":
            report_content = text
            text = ai_service.get_page_text(text)

        if mode == "document" and analysis_kind != "Image" and not text.strip():
            raise ValueError("No readable text was found in the document.")
        answer = ai_service.check_content(
            text, lang, analysis_kind,
            file_bytes if mode == "file" or (mode == "document" and mime) else None, mime
        )
        report = {"id": uuid.uuid4().hex[:10].upper(), "user_id": current_user_id(), "kind": report_kind,
                  "content": (report_content or report_kind)[:2000], "result": answer["result"], "confidence": answer["confidence"],
                  "reason": answer.get("reason", ""), "details": answer.get("details", ""), "sources": answer["sources"],
                  "created": datetime.now().strftime("%d %b %Y, %H:%M")}
        database.save_report(report)
        return jsonify(report)
    except Exception as error:
        return jsonify({"error": str(error)}), 400


# ---------- 4. REPORT, HISTORY, CHAT, PDF ----------
@app.route("/api/report/<report_id>")
def get_report(report_id):
    """Public: QR scan karne wala bina login ke report dekh sakta hai (sirf padhna). History/chat/PDF sirf maalik ke liye."""
    report = database.find_public_report(report_id)
    if not report:
        return jsonify({"error": "Report nahi mili"}), 404
    report["is_owner"] = report["user_id"] == current_user_id()   # kya ye isi user ki report hai?
    report.pop("user_id", None)                                   # user ki pehchaan bahar nahi jaati
    return jsonify(report)

@app.route("/api/history")
@login_required
def history():
    return jsonify(database.history_of(current_user_id()))

def my_report(report_id):
    """Report tabhi milegi jab wo isi user ki ho."""
    return database.find_report(report_id, current_user_id())

@app.route("/api/chat/<report_id>", methods=["POST"])
@login_required
def chat(report_id):
    report = my_report(report_id)
    if not report:
        return jsonify({"answer": "Report nahi mili. Login check karo."}), 404
    body = request.get_json(silent=True) or {}
    try:
        return jsonify({"answer": ai_service.chat_answer(report, str(body.get("q", "")), body.get("lang", "en"))})
    except Exception as error:
        return jsonify({"answer": f"Error: {error}"}), 500

@app.route("/api/report/<report_id>/pdf")
@login_required
def pdf(report_id):
    report = my_report(report_id)
    if not report:
        return "Report nahi mili. Login karo.", 404
    base_url = (os.getenv("PUBLIC_URL") or request.host_url).rstrip("/")
    return send_file(pdf_service.make_pdf(report, f"{base_url}/?id={report_id}"),
                     mimetype="application/pdf", download_name=f"VerifAI-{report_id}.pdf")


if __name__ == "__main__":
    app.run(host="0.0.0.0")
